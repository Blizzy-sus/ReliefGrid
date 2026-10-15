"""Helpers for building the ReliefGrid report with python-docx (styles, fields, sections, tables, code, figures)."""
import copy
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement

BODY_FONT = "Calibri"
CODE_FONT = "Consolas"
FOOTER_TEXT = "ReliefGrid • BCSE302L • October 2026"
HEADER_TEXT = "ReliefGrid — Disaster Relief Resource Allocation System"
PAGE_W = 8.5
PAGE_H = 11.0
MARGIN = 1.0
TEXT_W = PAGE_W - 2 * MARGIN  # 6.5 in


def set_font(run_or_style, name=BODY_FONT, size=None, bold=None, italic=None, color=None):
    f = run_or_style.font
    f.name = name
    rpr = run_or_style.element.get_or_add_rPr() if hasattr(run_or_style.element, "get_or_add_rPr") else None
    if rpr is not None:
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(a), name)
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = RGBColor.from_string(color)


def _para_border(pPr, side="bottom", sz=8, color="000000", space=2):
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        pPr.append(pbdr)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pbdr.append(el)


def new_document():
    doc = Document()
    st = doc.styles
    normal = st["Normal"]
    set_font(normal, BODY_FONT, 11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, before, after, italic in (("Heading 1", 16, 20, 8, False), ("Heading 2", 13, 12, 4, False), ("Heading 3", 11.5, 8, 3, True)):
        s = st[name]
        set_font(s, BODY_FONT, size, bold=True, italic=italic, color="000000")
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.line_spacing = 1.0
    _para_border(st["Heading 1"].element.get_or_add_pPr(), "bottom", 10, "000000", 1)

    cap = st["Caption"]
    set_font(cap, BODY_FONT, 9.5, bold=False, italic=True, color="404040")
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(10)

    # front-matter heading (not in the TOC outline)
    fh = st.add_style("Front Heading", WD_STYLE_TYPE.PARAGRAPH)
    fh.base_style = normal
    set_font(fh, BODY_FONT, 16, bold=True, color="000000")
    fh.paragraph_format.space_before = Pt(8)
    fh.paragraph_format.space_after = Pt(5)
    fh.paragraph_format.keep_with_next = True
    _para_border(fh.element.get_or_add_pPr(), "bottom", 10, "000000", 1)

    for n in ("TOC 1", "TOC 2", "Table of Figures"):
        try:
            s = st[n]
        except KeyError:
            s = st.add_style(n, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = normal
        set_font(s, BODY_FONT, (9 if n == "Table of Figures" else 9.5) if n != "TOC 1" else 10.5, bold=(n == "TOC 1"))
        s.paragraph_format.space_after = Pt(0.5 if n == "Table of Figures" else (1 if n != "TOC 1" else 2))
        s.paragraph_format.space_before = Pt(3 if n == "TOC 1" else 0)
        s.paragraph_format.line_spacing = 1.0
        if n == "TOC 2":
            s.paragraph_format.left_indent = Inches(0.25)

    # first section geometry
    sec = doc.sections[0]
    sec.page_width = Inches(PAGE_W)
    sec.page_height = Inches(PAGE_H)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(MARGIN))
    sec.header_distance = Inches(0.45)
    sec.footer_distance = Inches(0.45)
    return doc


# ------------------------------------------------------------------ fields
def add_field(paragraph, instr, placeholder="", bold=None, size=None, dirty=False):
    r = paragraph.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    if dirty:
        fc.set(qn("w:dirty"), "true")
    r._r.append(fc)
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r2._r.append(it)
    r3 = paragraph.add_run()
    fc2 = OxmlElement("w:fldChar")
    fc2.set(qn("w:fldCharType"), "separate")
    r3._r.append(fc2)
    r4 = paragraph.add_run(placeholder)
    r5 = paragraph.add_run()
    fc3 = OxmlElement("w:fldChar")
    fc3.set(qn("w:fldCharType"), "end")
    r5._r.append(fc3)
    for r_ in (r, r2, r3, r4, r5):
        if size:
            r_.font.size = Pt(size)
        if bold is not None:
            r_.font.bold = bold
    return r4


# ------------------------------------------------------------------ sections
def _sect_pr(section):
    return section._sectPr


def configure_section(section, landscape=False, border=True, num_fmt=None, num_start=None, show_header=True):
    sp = _sect_pr(section)
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(PAGE_H)
        section.page_height = Inches(PAGE_W)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(PAGE_W)
        section.page_height = Inches(PAGE_H)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(MARGIN if not landscape else 0.7))
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)
    for tag in ("w:pgBorders", "w:pgNumType"):
        for el in sp.findall(qn(tag)):
            sp.remove(el)
    pgmar = sp.find(qn("w:pgMar"))
    anchor = pgmar
    if border:
        b = parse_xml(
            f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page">'
            + "".join(f'<w:{s} w:val="single" w:sz="6" w:space="24" w:color="000000"/>' for s in ("top", "left", "bottom", "right"))
            + "</w:pgBorders>"
        )
        anchor.addnext(b)
        anchor = b
    if num_fmt or num_start:
        n = OxmlElement("w:pgNumType")
        if num_fmt:
            n.set(qn("w:fmt"), num_fmt)
        if num_start is not None:
            n.set(qn("w:start"), str(num_start))
        anchor.addnext(n)


def set_header_footer(section, with_page_number=True, with_header=True):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    for r in list(hp.runs):
        r._r.getparent().remove(r._r)
    if with_header:
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = hp.add_run(HEADER_TEXT)
        set_font(run, BODY_FONT, 9, italic=True, color="595959")
        _para_border(hp._p.get_or_add_pPr(), "bottom", 4, "A6A6A6", 3)
    fp = section.footer.paragraphs[0]
    for r in list(fp.runs):
        r._r.getparent().remove(r._r)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run(FOOTER_TEXT)
    set_font(run, BODY_FONT, 9, color="404040")
    if with_page_number:
        r2 = fp.add_run("   |   Page ")
        set_font(r2, BODY_FONT, 9, color="404040")
        add_field(fp, "PAGE", "1", size=9)


def new_section(doc, landscape=False, **kw):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(sec, landscape=landscape, **kw)
    return sec


# ------------------------------------------------------------------ text
def para(doc, text="", bold=False, italic=False, size=None, align=None, space_after=None, keep=False, color=None, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        add_runs(p, text, bold=bold, italic=italic, size=size, color=color)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def add_runs(p, text, bold=False, italic=False, size=None, color=None):
    """Supports **bold**, *italic* and `code` inline markup."""
    import re
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)", text)
    for t in tokens:
        if not t:
            continue
        if t.startswith("**") and t.endswith("**"):
            r = p.add_run(t[2:-2]); r.bold = True
        elif t.startswith("`") and t.endswith("`"):
            r = p.add_run(t[1:-1]); set_font(r, CODE_FONT, (size or 11) - 1.5)
            r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        elif t.startswith("*") and t.endswith("*") and len(t) > 2:
            r = p.add_run(t[1:-1]); r.italic = True
        else:
            r = p.add_run(t)
            if bold:
                r.bold = True
        if italic:
            r.italic = True
        if size and not t.startswith("`"):
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    return p


def h1(doc, text):
    return doc.add_heading(text, level=1)


def h2(doc, text):
    return doc.add_heading(text, level=2)


def h3(doc, text):
    return doc.add_heading(text, level=3)


def front_heading(doc, text):
    return doc.add_paragraph(text, style="Front Heading")


def bullets(doc, items, size=None, indent=0.25):
    for it in items:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Inches(indent + 0.2)
        pf.first_line_indent = Inches(-0.2)
        pf.space_after = Pt(3)
        r = p.add_run("•  ")
        if size:
            r.font.size = Pt(size)
        add_runs(p, it, size=size)


def numbered(doc, items, size=None, indent=0.25):
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Inches(indent + 0.28)
        pf.first_line_indent = Inches(-0.28)
        pf.space_after = Pt(3)
        r = p.add_run(f"{i}.  ")
        if size:
            r.font.size = Pt(size)
        add_runs(p, it, size=size)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------ tables
def _cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for el in tcPr.findall(qn("w:shd")):
        tcPr.remove(el)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _cell_margins(table, top=50, bottom=50, left=80, right=80):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    tblPr.append(m)


def _borders(table, color="A6A6A6", sz=4):
    tblPr = table._tbl.tblPr
    for el in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(el)
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        b.append(el)
    tblPr.append(b)


def _row_props(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trPr.append(cs)
    if header:
        th = OxmlElement("w:tblHeader")
        trPr.append(th)


def table(doc, headers, rows, widths, font=9.5, header_fill="E7E6E6", align=None, bold_first_col=False, zebra=False, total_width=None, keep=None, pad=50):
    """Create a bordered table. `align` is a list of 'l'/'r'/'c' per column."""
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borders(t)
    _cell_margins(t, pad, pad)
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    align = align or ["l"] * len(headers)

    def put(cell, text, bold=False, a="l", fill=None):
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}[a]
        lines = str(text).split("\n")
        for i, line in enumerate(lines):
            if i:
                p = cell.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}[a]
            add_runs(p, line, bold=bold, size=font)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if fill:
            _cell_shade(cell, fill)

    for i, h in enumerate(headers):
        put(t.rows[0].cells[i], h, bold=True, a="c" if align[i] == "c" else ("r" if align[i] == "r" else "l"), fill=header_fill)
    _row_props(t.rows[0], header=True)
    for ri, r in enumerate(rows):
        row = t.add_row()
        _row_props(row)
        for i, v in enumerate(r):
            put(row.cells[i], v, bold=(bold_first_col and i == 0), a=align[i], fill=("F7F7F7" if zebra and ri % 2 else None))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    # table grid widths
    grid = t._tbl.find(qn("w:tblGrid"))
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i] * 1440)))
    if keep is None:
        keep = len(rows) <= 7
    for ri, row in enumerate(t.rows):
        last = ri == len(t.rows) - 1
        if ri == 0 or (keep and not last):
            for c in row.cells:
                for pp in c.paragraphs:
                    pp.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def kv_table(doc, rows, widths=(2.4, 4.1), font=11):
    """Cover-page style two-column table (label | value), centred."""
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borders(t, "7F7F7F", 4)
    _cell_margins(t, 70, 70, 100, 100)
    for k, v in rows:
        row = t.add_row()
        for i, (txt, b) in enumerate(((k, True), (v, False))):
            c = row.cells[i]
            c.text = ""
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            add_runs(p, txt, bold=b, size=font)
            c.width = Inches(widths[i])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return t


# ------------------------------------------------------------------ code
def code_block(doc, text, font=8.5, fill="F3F4F6", keep_together=False):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borders(t, "BFBFBF", 4)
    _cell_margins(t, 70, 70, 120, 120)
    cell = t.rows[0].cells[0]
    cell.width = Inches(TEXT_W)
    _cell_shade(cell, fill)
    cell.text = ""
    lines = text.rstrip("\n").split("\n")
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        if keep_together and i < len(lines) - 1:
            p.paragraph_format.keep_with_next = True
        r = p.add_run(line if line else " ")
        set_font(r, CODE_FONT, font)
    if keep_together:
        _row_props(t.rows[0])
    grid = t._tbl.find(qn("w:tblGrid"))
    for gc in grid.findall(qn("w:gridCol")):
        gc.set(qn("w:w"), str(int(TEXT_W * 1440)))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def callout(doc, text, fill="EEF4FB", border="9DB7D5", size=10):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borders(t, border, 6)
    _cell_margins(t, 90, 90, 140, 140)
    c = t.rows[0].cells[0]
    c.width = Inches(TEXT_W)
    _cell_shade(c, fill)
    c.text = ""
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_runs(p, line, size=size)
    _row_props(t.rows[0])
    grid = t._tbl.find(qn("w:tblGrid"))
    for gc in grid.findall(qn("w:gridCol")):
        gc.set(qn("w:w"), str(int(TEXT_W * 1440)))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ------------------------------------------------------------------ captions & figures
def caption(doc, label, text):
    p = doc.add_paragraph(style="Caption")
    r = p.add_run(f"{label} ")
    r.bold = True
    add_field(p, f"SEQ {label} \\* ARABIC", "0", bold=True)
    r2 = p.add_run(f": {text}")
    return p


def figure(doc, path, caption_text, width=TEXT_W, max_height=8.0, label="Figure", keep_caption=True):
    from PIL import Image
    with Image.open(path) as im:
        w, h = im.size
    aspect = h / w
    width_in = min(width, max_height / aspect)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = keep_caption
    p.add_run().add_picture(path, width=Inches(width_in))
    caption(doc, label, caption_text)


def table_caption(doc, text):
    p = caption(doc, "Table", text)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p
